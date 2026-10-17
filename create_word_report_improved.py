#!/usr/bin/env python3
"""
Create Word document from Markdown report
Removes emojis and converts markdown lists to proper formatting without hyphens
"""

import re
import sys

try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
except ImportError:
    print("Installing python-docx...")
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx", "--quiet"])
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn

def remove_emojis(text):
    """Remove all emojis from text."""
    # Remove common emoji patterns
    emoji_pattern = re.compile(
        "["
        "\U0001F600-\U0001F64F"  # emoticons
        "\U0001F300-\U0001F5FF"  # symbols & pictographs
        "\U0001F680-\U0001F6FF"  # transport & map symbols
        "\U0001F1E0-\U0001F1FF"  # flags
        "\U00002702-\U000027B0"
        "\U000024C2-\U0001F251"
        "\U0001F900-\U0001F9FF"  # supplemental symbols
        "\U0001FA00-\U0001FA6F"  # chess symbols
        "\U0001FA70-\U0001FAFF"  # symbols and pictographs
        "\U00002600-\U000026FF"  # miscellaneous symbols
        "\U00002700-\U000027BF"  # dingbats
        "]+", flags=re.UNICODE)
    text = emoji_pattern.sub('', text)
    
    # Remove specific emoji characters
    emoji_chars = ['✅', '⚠️', '❌', '🚀', '📋', '📊', '📝', '🎯', '✨', '🔧', '📄']
    for emoji in emoji_chars:
        text = text.replace(emoji, '')
    
    return text.strip()

def clean_line(line):
    """Clean markdown line - remove list hyphens, emojis, and extra spaces."""
    # Remove emojis first
    line = remove_emojis(line)
    
    # Remove markdown list hyphens and asterisks at start
    line = re.sub(r'^[\s]*[-]\s+', '', line)
    line = re.sub(r'^[\s]*\*\s+', '', line)
    
    # Clean up multiple spaces
    line = re.sub(r'\s+', ' ', line)
    
    return line.strip()

def parse_markdown_to_doc(md_file, docx_file):
    """Parse markdown file and create Word document."""
    doc = Document()
    
    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    
    with open(md_file, 'r', encoding='utf-8') as f:
        content = f.read()
    
    # Split into lines
    lines = content.split('\n')
    
    i = 0
    in_code_block = False
    code_block_lines = []
    code_language = ''
    
    while i < len(lines):
        line = lines[i]
        original_line = line
        
        # Handle code blocks
        if line.strip().startswith('```'):
            if in_code_block:
                # End code block
                if code_block_lines:
                    para = doc.add_paragraph()
                    run = para.add_run('\n'.join(code_block_lines))
                    run.font.name = 'Courier New'
                    run.font.size = Pt(9)
                    para.style = 'No Spacing'
                code_block_lines = []
                in_code_block = False
                code_language = ''
            else:
                # Start code block
                in_code_block = True
                code_language = line.strip()[3:].strip()
            i += 1
            continue
        
        if in_code_block:
            code_block_lines.append(line)
            i += 1
            continue
        
        # Clean the line
        cleaned = clean_line(line)
        
        # Skip empty lines (but add spacing)
        if not cleaned and not original_line.strip():
            if i < len(lines) - 1:  # Don't add extra space at end
                doc.add_paragraph()
            i += 1
            continue
        
        # Handle headers
        if line.startswith('# '):
            heading = remove_emojis(cleaned)
            para = doc.add_heading(heading, level=1)
        elif line.startswith('## '):
            heading = remove_emojis(cleaned)
            para = doc.add_heading(heading, level=2)
        elif line.startswith('### '):
            heading = remove_emojis(cleaned)
            para = doc.add_heading(heading, level=3)
        elif line.startswith('#### '):
            heading = remove_emojis(cleaned)
            para = doc.add_heading(heading, level=4)
        # Handle horizontal rules
        elif line.strip() == '---' or line.strip() == '***':
            para = doc.add_paragraph('_' * 80)
        # Handle list items (convert to bullet points without hyphens)
        elif original_line.strip().startswith('-') or original_line.strip().startswith('*'):
            if cleaned:
                # Remove any remaining hyphens/asterisks
                cleaned = re.sub(r'^[-*]\s*', '', cleaned)
                cleaned = remove_emojis(cleaned)
                if cleaned:
                    para = doc.add_paragraph(cleaned, style='List Bullet')
        # Handle numbered lists
        elif re.match(r'^\d+\.\s+', original_line.strip()):
            if cleaned:
                # Extract number and text
                match = re.match(r'^(\d+)\.\s+(.*)', original_line.strip())
                if match:
                    num, text = match.groups()
                    text = remove_emojis(text)
                    para = doc.add_paragraph(text, style='List Number')
        # Handle bold text
        elif '**' in line:
            para = doc.add_paragraph()
            parts = re.split(r'(\*\*[^*]+\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = para.add_run(part[2:-2])
                    run.bold = True
                elif part.strip():
                    run = para.add_run(remove_emojis(part))
        # Handle regular paragraphs
        else:
            if cleaned:
                cleaned = remove_emojis(cleaned)
                # Remove any remaining markdown formatting
                cleaned = re.sub(r'\*\*([^*]+)\*\*', r'\1', cleaned)  # Remove bold
                cleaned = re.sub(r'`([^`]+)`', r'\1', cleaned)  # Remove code inline
                if cleaned:
                    para = doc.add_paragraph(cleaned)
        
        i += 1
    
    # Save document
    doc.save(docx_file)
    print(f"Word document created successfully: {docx_file}")

def main():
    md_file = 'MLOPS_ASSIGNMENT_REPORT.md'
    docx_file = 'MLOPS_ASSIGNMENT_REPORT.docx'
    
    try:
        parse_markdown_to_doc(md_file, docx_file)
        print(f"Successfully created Word document: {docx_file}")
        print("All emojis removed and list hyphens converted to proper formatting.")
    except Exception as e:
        print(f"Error creating Word document: {e}")
        import traceback
        traceback.print_exc()
        sys.exit(1)

if __name__ == '__main__':
    main()

